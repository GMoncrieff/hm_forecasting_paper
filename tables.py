"""Table-generating functions. Outputs to config.OUTPUT_DIR.

Each table is rendered as a reportlab PDF, a python-docx DOCX, or both,
according to config.TABLE_FORMAT. The two renderers consume the same loaded
data; only the layout code differs.
"""
from xml.sax.saxutils import escape

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

import config

# --- output-format dispatch -------------------------------------------------

VALID_TABLE_FORMATS = ('pdf', 'docx', 'both')

RULE_HEAVY = 8   # eighths of a point -> 1.0pt, matches the reportlab rules
RULE_LIGHT = 4   # eighths of a point -> 0.5pt
GROUP_FILL = 'EFEFEF'
LINK_COLOR = '1A56A0'


def _formats() -> tuple[str, ...]:
    """Return the formats to emit, per config.TABLE_FORMAT."""
    fmt = str(getattr(config, 'TABLE_FORMAT', 'pdf')).strip().lower()
    if fmt not in VALID_TABLE_FORMATS:
        raise ValueError(
            f'config.TABLE_FORMAT is {fmt!r}; expected one of '
            f'{", ".join(VALID_TABLE_FORMATS)}'
        )
    return ('pdf', 'docx') if fmt == 'both' else (fmt,)


# --- python-docx helpers ----------------------------------------------------
# python-docx is imported lazily so that PDF-only runs need no extra dependency.

# Schema-mandated child order for <w:tcPr>; Word rejects other orderings.
_TCPR_ORDER = (
    'w:cnfStyle', 'w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge', 'w:tcBorders',
    'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign',
    'w:hideMark',
)


def _require_docx():
    """Return the python-docx module, with an actionable error if it is absent."""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise ImportError(
            f'config.TABLE_FORMAT is {config.TABLE_FORMAT!r}, which requests DOCX '
            'output, but python-docx is not installed. Install it with:\n'
            '    mamba install -n hm_plots -c conda-forge python-docx'
        ) from exc
    return docx


def _tcpr_set(cell, element, tag: str):
    """Insert `element` into the cell's <w:tcPr> at its schema position."""
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn(tag))
    if existing is not None:
        tc_pr.remove(existing)
    later = {qn(t) for t in _TCPR_ORDER[_TCPR_ORDER.index(tag) + 1:]}
    for child in tc_pr:
        if child.tag in later:
            child.addprevious(element)
            return element
    tc_pr.append(element)
    return element


def _new_docx(title: str):
    """A landscape-letter document, 0.5in margins, Times New Roman, no padding."""
    docx = _require_docx()
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = docx.Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    # python-docx does not swap the page dimensions when orientation changes.
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(section, f'{side}_margin', Inches(0.5))

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    # The east-Asian face must be named separately or Word substitutes a default.
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn('w:eastAsia'), 'Times New Roman')
    # Word's stock 8pt paragraph spacing would treble the height of every row.
    fmt = normal.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0

    doc.core_properties.title = title
    return doc


def _new_table(doc, n_cols: int, widths_in):
    """An unruled table with fixed column widths (Word ignores them otherwise)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    table = doc.add_table(rows=0, cols=n_cols)
    table.autofit = False

    # Start rule-free regardless of the template default; the paper style uses
    # horizontal rules only, added per-cell below.
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        borders.append(el)
    table._tbl.tblPr.append(borders)

    for column, width in zip(table.columns, widths_in):
        column.width = Inches(width)
    return table


def _add_row(table, widths_in):
    """Append a row, applying the per-cell widths Word needs to honour a layout."""
    from docx.shared import Inches

    row = table.add_row()
    for cell, width in zip(row.cells, widths_in):
        cell.width = Inches(width)
    return row


def _rule(cell, edge: str, size: int = RULE_HEAVY):
    """Draw a horizontal rule on one edge ('top' or 'bottom') of a cell."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn('w:tcBorders'))
    if borders is None:
        borders = _tcpr_set(cell, OxmlElement('w:tcBorders'), 'w:tcBorders')
    existing = borders.find(qn(f'w:{edge}'))
    if existing is not None:
        borders.remove(existing)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(size))
    el.set(qn('w:color'), '000000')
    borders.append(el)


def _shade(cell, fill: str = GROUP_FILL):
    """Fill a cell with a solid background colour."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill.lstrip('#'))
    _tcpr_set(cell, shd, 'w:shd')


def _repeat_header(row):
    """Repeat this row atop each new page - the analogue of reportlab repeatRows."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    row._tr.get_or_add_trPr().append(el)


def _write(paragraph, text, *, bold=False, size=8.5, color=None, strip=True):
    """Write text into a paragraph, rendering any '⁻¹' as a superscript run.

    Mirrors the <super>-1</super> substitution the PDF path applies in markup().
    Cell values carry stray whitespace from Excel, so they are stripped by
    default; pass strip=False for text whose spacing is deliberate.
    """
    from docx.shared import Pt, RGBColor

    def run(value, superscript=False):
        r = paragraph.add_run(value)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.superscript = superscript
        if color is not None:
            r.font.color.rgb = RGBColor.from_string(color)

    text = '' if pd.isna(text) else str(text)
    if strip:
        text = text.strip()
    for i, part in enumerate(text.split('⁻¹')):
        if i:                      # a '⁻¹' separator was consumed by the split
            run('-1', superscript=True)
        if part:
            run(part)


def _fill(cell, text, *, bold=False, size=8.5, align='left', valign='top'):
    """Set a cell's sole paragraph to `text` with the given alignment."""
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    cell.vertical_alignment = {
        'top': WD_ALIGN_VERTICAL.TOP, 'middle': WD_ALIGN_VERTICAL.CENTER,
    }[valign]
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {
        'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
    }[align]
    _write(paragraph, text, bold=bold, size=size)
    return paragraph


def _fill_link(cell, url: str, *, size=8.5):
    """Set a cell to an external hyperlink; python-docx has no API for these."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = cell.paragraphs[0]
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'),
             paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))

    run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    for tag, value in (('w:color', LINK_COLOR), ('w:sz', str(int(size * 2)))):
        el = OxmlElement(tag)          # <w:color> must precede <w:sz> per schema
        el.set(qn('w:val'), value)
        r_pr.append(el)
    run.append(r_pr)

    text = OxmlElement('w:t')
    text.text = url
    text.set(qn('xml:space'), 'preserve')
    run.append(text)

    link.append(run)
    paragraph._p.append(link)


def _merge_across(row, text, *, size=9.0):
    """Turn a row into a shaded, bold heading spanning the full table width."""
    cells = row.cells
    merged = cells[0].merge(cells[-1])
    merged.text = ''
    _fill(merged, text, bold=True, size=size, align='left', valign='middle')
    _shade(merged)
    return merged


def table_1():
    """Table 1: model covariates -> landscape PDF and/or DOCX.

    Main-text table, not supplementary. The source workbook's own title row
    still reads "Supplementary Table S2"; it is not rendered (only the
    description and notes rows are), but it is worth correcting at source.
    """
    OUT = config.OUTPUT_DIR
    OUT.mkdir(parents=True, exist_ok=True)
    XLSX = config.PATHS['covariates_xlsx']
    OUT_PDF = OUT / 'Table_1_covariates.pdf'
    OUT_DOCX = OUT_PDF.with_suffix('.docx')
    DOC_TITLE = 'Table 1. Model covariates'

    TABLE_LABEL = 'Table 1'

    # Columns to show (the source `Type` column is consumed as a group subheading).
    DISPLAY_COLS = [
        'Covariate', 'Source dataset', 'Native spatial resolution',
        'Native temporal resolution', 'Time period used', 'Units',
        'Citation', 'Link / DOI',
    ]
    COL_WIDTHS = [1.55, 2.05, 0.95, 1.05, 1.15, 0.80, 1.05, 1.40]  # inches, sum=10.0
    LINK_COL = DISPLAY_COLS.index('Link / DOI')

    # Order groups appear in the table; label shown on the spanning subheading row.
    GROUP_LABEL = {
        'Static': 'Static covariates',
        'Dynamic': 'Dynamic covariates',
        'Static (learned)': 'Static (learned) covariates',
    }

    CELL = ParagraphStyle('cell', fontName='Times-Roman', fontSize=8.5, leading=10)
    HEAD = ParagraphStyle('head', fontName='Times-Bold', fontSize=8.5, leading=10)
    LINK = ParagraphStyle('link', fontName='Times-Roman', fontSize=8.5, leading=10,
                          textColor=colors.HexColor('#1A56A0'))


    def markup(text: str) -> str:
        """XML-escape a cell value, then restore the markup we intend to keep."""
        s = escape('' if pd.isna(text) else str(text).strip())
        s = s.replace('⁻¹', '<super>-1</super>')
        return s


    def load() -> tuple[str, pd.DataFrame, str]:
        """Return (description, data-with-Type, notes) parsed from the workbook."""
        raw = pd.read_excel(XLSX, header=None)
        col0 = raw.iloc[:, 0].astype('string')
        hdr = col0.eq('Covariate').idxmax()                 # header row index
        description = str(raw.iloc[hdr - 2, 0]).strip()     # row above the blank
        notes_idx = col0.str.startswith('Notes', na=False).idxmax()
        notes = str(raw.iloc[notes_idx, 0]).strip()

        df = raw.iloc[hdr + 1:notes_idx].copy()
        df.columns = raw.iloc[hdr].tolist()
        df = df.dropna(how='all').reset_index(drop=True)
        return description, df, notes


    def build_table(df: pd.DataFrame) -> Table:
        header = [Paragraph(markup(c), HEAD) for c in DISPLAY_COLS]
        rows = [header]

        style = [
            ('LINEABOVE',     (0, 0), (-1, 0), 1.0, colors.black),
            ('LINEBELOW',     (0, 0), (-1, 0), 1.0, colors.black),
            ('VALIGN',        (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING',   (0, 0), (-1, -1), 4),
            ('RIGHTPADDING',  (0, 0), (-1, -1), 4),
            ('TOPPADDING',    (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]

        group_rows = []
        for gtype in GROUP_LABEL:
            grp = df[df['Type'].astype('string').str.strip() == gtype]
            if grp.empty:
                continue
            rows.append([GROUP_LABEL[gtype], '', '', '', '', '', '', ''])
            group_rows.append(len(rows) - 1)
            for _, r in grp.iterrows():
                cells = []
                for j, col in enumerate(DISPLAY_COLS):
                    if j == LINK_COL and not pd.isna(r[col]):
                        url = str(r[col]).strip()
                        cells.append(Paragraph(
                            f'<link href="{escape(url)}">{escape(url)}</link>', LINK))
                    else:
                        cells.append(Paragraph(markup(r[col]), CELL))
                rows.append(cells)

        for ri in group_rows:
            style += [
                ('SPAN',          (0, ri), (-1, ri)),
                ('FONTNAME',      (0, ri), (-1, ri), 'Times-Bold'),
                ('FONTSIZE',      (0, ri), (-1, ri), 9),
                ('ALIGN',         (0, ri), (-1, ri), 'LEFT'),
                ('BACKGROUND',    (0, ri), (-1, ri), colors.HexColor('#EFEFEF')),
                ('TOPPADDING',    (0, ri), (-1, ri), 4),
                ('BOTTOMPADDING', (0, ri), (-1, ri), 4),
            ]
        style.append(('LINEBELOW', (0, -1), (-1, -1), 1.0, colors.black))

        table = Table(rows, colWidths=[w * inch for w in COL_WIDTHS], repeatRows=1)
        table.setStyle(TableStyle(style))
        return table


    def build_docx(description: str, df: pd.DataFrame, notes: str) -> None:
        _require_docx()      # actionable error before any submodule import
        from docx.shared import Pt

        doc = _new_docx(DOC_TITLE)

        caption = doc.add_paragraph()
        caption.paragraph_format.space_after = Pt(8)
        _write(caption, f'{TABLE_LABEL}.', bold=True, size=11)
        _write(caption, f' {description}', size=11, strip=False)

        table = _new_table(doc, len(DISPLAY_COLS), COL_WIDTHS)

        header = _add_row(table, COL_WIDTHS)
        for cell, name in zip(header.cells, DISPLAY_COLS):
            _fill(cell, name, bold=True)
            _rule(cell, 'top')
            _rule(cell, 'bottom')
        _repeat_header(header)

        last = header
        for gtype, label in GROUP_LABEL.items():
            grp = df[df['Type'].astype('string').str.strip() == gtype]
            if grp.empty:
                continue
            _merge_across(_add_row(table, COL_WIDTHS), label)
            for _, r in grp.iterrows():
                last = _add_row(table, COL_WIDTHS)
                for j, col in enumerate(DISPLAY_COLS):
                    if j == LINK_COL and not pd.isna(r[col]):
                        _fill_link(last.cells[j], str(r[col]).strip())
                    else:
                        _fill(last.cells[j], r[col])
        for cell in last.cells:
            _rule(cell, 'bottom')

        footer = doc.add_paragraph()
        footer.paragraph_format.space_before = Pt(8)
        _write(footer, notes, size=8)

        doc.save(str(OUT_DOCX))
        print(f'Saved {OUT_DOCX}  ({len(df)} covariates)')

    def build_pdf(description: str, df: pd.DataFrame, notes: str) -> None:
        doc = SimpleDocTemplate(
            str(OUT_PDF), pagesize=landscape(letter),
            leftMargin=0.5 * inch, rightMargin=0.5 * inch,
            topMargin=0.5 * inch, bottomMargin=0.5 * inch,
            title=DOC_TITLE,
        )

        caption_style = ParagraphStyle(
            'caption', fontName='Times-Roman', fontSize=11, leading=14, spaceAfter=8)
        notes_style = ParagraphStyle(
            'notes', fontName='Times-Roman', fontSize=8, leading=10, spaceBefore=8)

        story = [
            Paragraph(f'<b>{TABLE_LABEL}.</b> {markup(description)}', caption_style),
            build_table(df),
            Spacer(1, 4),
            Paragraph(markup(notes), notes_style),
        ]
        doc.build(story)
        print(f'Saved {OUT_PDF}  ({len(df)} covariates)')

    def main() -> None:
        formats = _formats()
        description, df, notes = load()
        if 'pdf' in formats:
            build_pdf(description, df, notes)
        if 'docx' in formats:
            build_docx(description, df, notes)

    main()


def tables_s1_s8():
    """Supplementary Tables S1-S8: per-realm ecoregion composition -> PDF/DOCX.

    One table per realm in REALM_ORDER, so Afrotropic is S1 and Palearctic S8.

    Depends on fig10's unprotected_loss_stats.csv; computes it (via fig10) if absent.
    """
    FORMATS = _formats()   # validate before a missing stats CSV triggers fig10

    OUT = config.OUTPUT_DIR
    OUT.mkdir(parents=True, exist_ok=True)
    STATS_CSV = OUT / 'unprotected_loss_stats.csv'
    OUT_DIR = OUT / 'realm_tables'
    TABLES_DIR = OUT_DIR
    OUT_PDF = OUT / 'realm_tables.pdf'
    OUT_DOCX = OUT / 'realm_tables.docx'
    DOC_TITLE = 'Per-realm ecoregion composition tables'
    if not STATS_CSV.exists():
        import figures
        figures.fig10()

    COLUMN_MAP = {
        'BIOME_NAME': 'Biome',
        'ECO_NAME': 'Ecoregion',
        'pct_protected': 'Protected',
        'pct_grey_upper': 'Still Natural 2040 (upper)',
        'pct_grey_central': 'Still Natural 2040 (central)',
        'pct_lost_upper': 'Natural lands loss 2040 (upper)',
        'pct_lost_central': 'Natural Lands loss 2040 (central)',
        'pct_unprot_nonnatural_2020': 'Non-natural 2020',
    }

    PROPORTION_COLS = [
        'Protected',
        'Still Natural 2040 (upper)',
        'Still Natural 2040 (central)',
        'Natural lands loss 2040 (upper)',
        'Natural Lands loss 2040 (central)',
        'Non-natural 2020',
    ]


    def _export_main() -> None:
        if not STATS_CSV.exists():
            raise SystemExit(
                f"Stats CSV not found: {STATS_CSV}\n"
                "Run plot_unprotected_loss.py first to generate it."
            )

        OUT_DIR.mkdir(parents=True, exist_ok=True)
        print(f"Reading {STATS_CSV.name}…")
        stats = pd.read_csv(STATS_CSV)

        missing = [c for c in COLUMN_MAP if c not in stats.columns]
        if missing:
            raise SystemExit(
                f"Stats CSV is missing required columns: {missing}\n"
                "Re-run plot_unprotected_loss.py to regenerate with the latest schema."
            )

        print(f"  {len(stats)} ecoregions across {stats['REALM'].nunique()} realms")

        for realm, group in stats.groupby('REALM'):
            out = group[list(COLUMN_MAP.keys())].rename(columns=COLUMN_MAP).copy()
            out[PROPORTION_COLS] = (out[PROPORTION_COLS] / 100.0).round(4)
            out = out.sort_values(['Biome', 'Ecoregion']).reset_index(drop=True)

            fname = f"{realm.replace(' ', '_').replace('/', '_')}.csv"
            out_path = OUT_DIR / fname
            out.to_csv(out_path, index=False)
            print(f"  saved {fname}  ({len(out)} rows)")

        print("=== done ===")

    REALM_ORDER = [
        'Afrotropic', 'Antarctica', 'Australasia', 'Indomalayan',
        'Nearctic', 'Neotropic', 'Oceania', 'Palearctic',
    ]

    REALM_ADJECTIVE = {
        'Afrotropic':  'Afrotropical',
        'Antarctica':  'Antarctic',
        'Australasia': 'Australasian',
        'Indomalayan': 'Indomalayan',
        'Nearctic':    'Nearctic',
        'Neotropic':   'Neotropical',
        'Oceania':     'Oceanian',
        'Palearctic':  'Palearctic',
    }

    # The protected mask is hm_static_iucn_strict_1000.tif — IUCN categories I–IV
    # are the conventional "strict" set. Adjust the label if your raster differs.
    IUCN_TEXT = 'IUCN categories I–IV'

    DECIMALS = 3

    # Column widths in inches; shared by both renderers.
    COL_WIDTHS = [
        3.10,  # Ecoregion (long names wrap)
        0.80,  # Protected
        0.95,  # Still Natural upper
        0.95,  # Still Natural central
        1.10,  # Natural loss upper (group label needs ~2.2in across cols 4-5)
        1.10,  # Natural loss central
        1.30,  # Non-natural 2020
    ]

    # Two-row header:
    # Row 0 — grouped labels only (over cols 2-3 and 4-5), other cols empty.
    # Row 1 — actual column titles for every column.
    # This keeps the underline below the grouped labels well above the
    # single-row column titles, so nothing is bisected by a line.
    HEADER_TOP = ['', '',
                  'Still Natural 2040', '',
                  'Natural lands loss 2040', '',
                  '']
    HEADER_BOTTOM = ['Ecoregion', 'Protected',
                     'upper', 'central',
                     'upper', 'central',
                     'Non-natural 2020']

    # PROPORTION_COLS is already in the order HEADER_BOTTOM names them, so it
    # doubles as the value-column order for both renderers.
    VALUE_COLS = PROPORTION_COLS


    def caption_parts(table_num: int, realm: str) -> tuple[str, str]:
        """(bold label, body) — one caption source for both renderers.

        The denominator is stated explicitly: the old wording ("proportion of
        {realm} ecoregions formally protected") read as a count of ecoregions
        rather than a share of each ecoregion's area, and said nothing about
        the grid the share was measured on.
        """
        grid_note = (
            f" Proportions are of ground area, measured on a "
            f"{config.EQUAL_AREA_RES // 1000} km equal-area grid "
            f"({config.EQUAL_AREA_CRS})."
        ) if config.EQUAL_AREA_CRS else ''
        return (
            # table_num counts from 1, and the realm tables now start at S1
            # (the covariates table moved to the main text as Table 1).
            f"Table S{table_num}.",
            f" Proportion of each {REALM_ADJECTIVE[realm]} ecoregion's land "
            f"area formally protected ({IUCN_TEXT}) in 2020, and the projected "
            f"2040 status of unprotected land under the upper and central "
            f"scenarios.{grid_note}"
        )


    def caption_html(table_num: int, realm: str) -> str:
        label, body = caption_parts(table_num, realm)
        return f"<b>{label}</b>{body}"


    def fmt(x: float) -> str:
        return f"{x:.{DECIMALS}f}"


    ECO_STYLE = ParagraphStyle(
        'eco_cell', fontName='Times-Roman', fontSize=12, leading=14,
    )


    def build_table(df: pd.DataFrame) -> Table:
        """Build a reportlab Table for one realm's ecoregion data."""
        df = df.sort_values(['Biome', 'Ecoregion']).reset_index(drop=True)

        rows = [list(HEADER_TOP), list(HEADER_BOTTOM)]

        style_cmds: list = [
            # Grouped header spans + underlines (only over the grouped cells)
            ('SPAN',      (2, 0), (3, 0)),
            ('SPAN',      (4, 0), (5, 0)),
            ('LINEBELOW', (2, 0), (3, 0), 0.5, colors.black),
            ('LINEBELOW', (4, 0), (5, 0), 0.5, colors.black),
            # Outer header rules
            ('LINEABOVE', (0, 0), (-1, 0), 1.0, colors.black),
            ('LINEBELOW', (0, 1), (-1, 1), 1.0, colors.black),
            # Fonts
            ('FONTNAME', (0, 0), (-1, 1), 'Times-Bold'),
            ('FONTNAME', (0, 2), (-1, -1), 'Times-Roman'),
            # Sizes / alignment
            ('FONTSIZE', (0, 0), (-1, -1), 12),
            ('LEADING',  (0, 0), (-1, -1), 14),
            ('VALIGN',   (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN',    (0, 0), (0, -1), 'LEFT'),
            ('ALIGN',    (1, 0), (-1, -1), 'CENTER'),
            # Padding
            ('LEFTPADDING',   (0, 0), (-1, -1), 4),
            ('RIGHTPADDING',  (0, 0), (-1, -1), 4),
            ('TOPPADDING',    (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]

        biome_row_indices = []
        for biome_name, biome_grp in df.groupby('Biome', sort=False):
            rows.append([biome_name, '', '', '', '', '', ''])
            biome_row_indices.append(len(rows) - 1)
            for _, r in biome_grp.iterrows():
                rows.append(
                    # Paragraph wraps long ecoregion names
                    [Paragraph(r['Ecoregion'], ECO_STYLE)]
                    + [fmt(r[c]) for c in VALUE_COLS]
                )

        # Bold biome subheading rows that span the whole table.
        for ri in biome_row_indices:
            style_cmds += [
                ('SPAN',     (0, ri), (-1, ri)),
                ('FONTNAME', (0, ri), (-1, ri), 'Times-Bold'),
                ('ALIGN',    (0, ri), (-1, ri), 'LEFT'),
                ('BACKGROUND', (0, ri), (-1, ri), colors.HexColor('#EFEFEF')),
                ('TOPPADDING',    (0, ri), (-1, ri), 4),
                ('BOTTOMPADDING', (0, ri), (-1, ri), 4),
            ]

        style_cmds.append(('LINEBELOW', (0, -1), (-1, -1), 1.0, colors.black))

        table = Table(rows, colWidths=[w * inch for w in COL_WIDTHS], repeatRows=2)
        table.setStyle(TableStyle(style_cmds))
        return table


    def build_docx_table(doc, df: pd.DataFrame) -> None:
        """Append one realm's ecoregion table to an open python-docx document."""
        df = df.sort_values(['Biome', 'Ecoregion']).reset_index(drop=True)

        table = _new_table(doc, len(HEADER_BOTTOM), COL_WIDTHS)
        top = _add_row(table, COL_WIDTHS)
        bottom = _add_row(table, COL_WIDTHS)
        for row in (top, bottom):
            _repeat_header(row)

        # Row 0 — outer top rule across everything, then the two grouped labels,
        # each underlined only across the columns it spans.
        for cell in top.cells:
            _rule(cell, 'top')
        for lo, hi in ((2, 3), (4, 5)):
            merged = top.cells[lo].merge(top.cells[hi])
            merged.text = ''
            _fill(merged, HEADER_TOP[lo], bold=True, size=12,
                  align='center', valign='middle')
            _rule(merged, 'bottom', RULE_LIGHT)

        # Row 1 — column titles, closed by the heavy rule under the header block.
        for j, (cell, title) in enumerate(zip(bottom.cells, HEADER_BOTTOM)):
            _fill(cell, title, bold=True, size=12,
                  align='left' if j == 0 else 'center', valign='middle')
            _rule(cell, 'bottom')

        last = bottom
        for biome_name, biome_grp in df.groupby('Biome', sort=False):
            _merge_across(_add_row(table, COL_WIDTHS), biome_name, size=12)
            for _, r in biome_grp.iterrows():
                last = _add_row(table, COL_WIDTHS)
                _fill(last.cells[0], r['Ecoregion'], size=12, valign='middle')
                for j, col in enumerate(VALUE_COLS, start=1):
                    _fill(last.cells[j], fmt(r[col]), size=12,
                          align='center', valign='middle')

        for cell in last.cells:
            _rule(cell, 'bottom')


    def _build_docx_main() -> None:
        if not TABLES_DIR.exists():
            raise SystemExit(f"Tables directory not found: {TABLES_DIR}")

        _require_docx()      # actionable error before any submodule import
        from docx.shared import Pt

        doc = _new_docx(DOC_TITLE)

        table_num = 0
        for realm in REALM_ORDER:
            csv_path = TABLES_DIR / f"{realm}.csv"
            if not csv_path.exists():
                print(f"  warning: {csv_path.name} not found, skipping")
                continue
            df = pd.read_csv(csv_path)
            table_num += 1
            if table_num > 1:
                # Break *before* each table but the first, so a missing realm
                # cannot leave a trailing blank page.
                doc.add_page_break()

            caption = doc.add_paragraph()
            caption.paragraph_format.space_after = Pt(8)
            label, body = caption_parts(table_num, realm)
            _write(caption, label, bold=True, size=12)
            _write(caption, body, size=12, strip=False)

            build_docx_table(doc, df)
            print(f"  Table {table_num}: {realm} ({len(df)} ecoregions)")

        doc.save(str(OUT_DOCX))
        print(f"Saved {OUT_DOCX}")


    def _build_main() -> None:
        if not TABLES_DIR.exists():
            raise SystemExit(f"Tables directory not found: {TABLES_DIR}")

        doc = SimpleDocTemplate(
            str(OUT_PDF),
            pagesize=landscape(letter),
            leftMargin=0.5 * inch,
            rightMargin=0.5 * inch,
            topMargin=0.5 * inch,
            bottomMargin=0.5 * inch,
            title='Per-realm ecoregion composition tables',
        )

        caption_style = ParagraphStyle(
            'caption',
            fontName='Times-Roman',
            fontSize=12,
            leading=15,
            spaceAfter=8,
        )

        story = []
        table_num = 0
        for realm in REALM_ORDER:
            csv_path = TABLES_DIR / f"{realm}.csv"
            if not csv_path.exists():
                print(f"  warning: {csv_path.name} not found, skipping")
                continue
            df = pd.read_csv(csv_path)
            table_num += 1
            story.append(Paragraph(caption_html(table_num, realm), caption_style))
            story.append(build_table(df))
            if realm != REALM_ORDER[-1]:
                story.append(PageBreak())
            print(f"  Table {table_num}: {realm} ({len(df)} ecoregions)")

        doc.build(story)
        print(f"Saved {OUT_PDF}")

    _export_main()
    if 'pdf' in FORMATS:
        _build_main()
    if 'docx' in FORMATS:
        _build_docx_main()
